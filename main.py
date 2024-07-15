import openpyxl
from docxtpl import DocxTemplate
import openpyxl.workbook
from docx import Document
from docxcompose.composer import Composer
from docx.shared import Inches
import os
from colorama import init, Fore, Style

# path excel workbook
excel_ws ='C:/Users/feder/Desktop/UNIONE_SCHEDE_BBCC_INV_REGIONE.xlsx'
word_template_path='C:/Users/feder/Desktop/Sample.docx'

out_path='C:/Users/feder/Desktop/Unione_Resume.docx'

img_path='C:/Users/feder/Desktop/Immagini/'

dummy='C:/Users/feder/Desktop/dummy.svg.png'

workbook = openpyxl.load_workbook(excel_ws)

workbook_sheet = workbook["Final_Data_Cini"]

list_sh=list(workbook_sheet.values)


main_doc = Document(word_template_path)
composer=Composer(main_doc)
header = list_sh[0]
init(autoreset=True)

print(Fore.YELLOW+"Processing....")

for index,row in enumerate(list_sh[1:], start=1):
    # context=dict(zip(header,row))
    context = {header[i]: (cell if cell is not None else '-')for i,cell in enumerate(row)}

    image_name=context.get('IMG_PATH','')
    
    # img_path_rel=f'{img_path}{image_name}'
    img_path_rel=os.path.join(img_path,image_name)


    if not os.path.isfile(img_path_rel):
        print(Fore.LIGHTRED_EX+Style.BRIGHT+'ERROR: '+Fore.BLUE+ f'Missing Cell Value or Incorrect Path for ROW: {Fore.GREEN}{index}')
        img_path_rel=dummy

    doc=DocxTemplate(word_template_path)
    doc.render(context)
    
    temp_out_P =f'temp.docx'
    doc.save(temp_out_P)

    sub_doc = Document(temp_out_P)

    for rel in sub_doc.part.rels.values():
        if "image" in rel.target_ref:
            rel.target_part._blob=open(img_path_rel, "rb").read()
            break

    # composer=Composer(main_doc)
    if index == 1:
        main_doc.add_page_break()
    composer.append(sub_doc)

    if row != list_sh[-1]:
        main_doc.add_page_break()
    


main_doc.save(out_path)

print(Fore.GREEN+Style.BRIGHT+'[SUCCESS]: '+ Fore.GREEN+"The file has been saved in " + out_path)